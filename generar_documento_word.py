import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, title, body, border_color="10B981", bg_color="ECFDF5", text_color="065F46"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(title + "\n")
    r1.bold = True
    r1.font.size = Pt(10.5)
    r1.font.name = "Segoe UI"
    r1.font.color.rgb = RGBColor(int(text_color[:2], 16), int(text_color[2:4], 16), int(text_color[4:], 16))

    r2 = p.add_run(body)
    r2.font.size = Pt(9.5)
    r2.font.name = "Segoe UI"
    r2.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = RGBColor(5, 150, 105) # Emerald Green Dark
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    return h

def add_heading_3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = RGBColor(6, 95, 70) # Dark Teal
    return h

def add_paragraph(doc, text, bold_prefix="", space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Segoe UI"
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    run = p.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(doc, text, bold_title=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_title:
        r_title = p.add_run(bold_title + " ")
        r_title.bold = True
        r_title.font.name = "Segoe UI"
        r_title.font.size = Pt(10)
        r_title.font.color.rgb = RGBColor(15, 23, 42)
    run = p.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_image_with_caption(doc, image_path, caption, width_inches=6.0):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(image_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(f"Figura: {caption}")
        r_cap.font.name = "Segoe UI"
        r_cap.font.size = Pt(9)
        r_cap.italic = True
        r_cap.font.color.rgb = RGBColor(100, 116, 139)
    else:
        add_paragraph(doc, f"[Imagen no encontrada: {os.path.basename(image_path)}]", space_after=8)

def create_document():
    doc = docx.Document()
    
    # Configurar márgenes de 2.5 cm (aproximadamente 1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # -------------------------------------------------------------
    # 1. PORTADA OFICIAL (20 PUNTOS)
    # -------------------------------------------------------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("UNIVERSIDAD TECNOLÓGICA METROPOLITANA")
    r_inst.bold = True
    r_inst.font.name = "Segoe UI"
    r_inst.font.size = Pt(18)
    r_inst.font.color.rgb = RGBColor(5, 150, 105)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(2)
    r_div = p_div.add_run("DIVISIÓN DE TECNOLOGÍAS DE LA INFORMACIÓN Y COMUNICACIÓN (TIC)")
    r_div.bold = True
    r_div.font.name = "Segoe UI"
    r_div.font.size = Pt(11)
    r_div.font.color.rgb = RGBColor(30, 41, 59)

    p_carr = doc.add_paragraph()
    p_carr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_carr.paragraph_format.space_before = Pt(0)
    p_carr.paragraph_format.space_after = Pt(12)
    r_carr = p_carr.add_run("Licenciatura en Ingeniería en Tecnologías de la Información e Innovación Digital\nÁrea Desarrollo de Software Multiplataforma (DSM)")
    r_carr.font.name = "Segoe UI"
    r_carr.font.size = Pt(10)
    r_carr.font.color.rgb = RGBColor(71, 85, 105)

    # Imagen de Portada (Banner temático generado)
    img_portada = r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Evidencias_Capturas\portada_banner.jpg"
    if os.path.exists(img_portada):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(12)
        p_img.add_run().add_picture(img_portada, width=Inches(5.0))

    # Título del Proyecto
    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_before = Pt(6)
    p_tit.paragraph_format.space_after = Pt(4)
    r_tit = p_tit.add_run("PRÁCTICA: ARREGLOS Y MATRICES EN C#")
    r_tit.bold = True
    r_tit.font.name = "Segoe UI"
    r_tit.font.size = Pt(20)
    r_tit.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Instrumento de Evaluación F-SGC-033 | Revisión 00\nUnidad I: Conceptos básicos de estructuras de datos orientadas a objetos")
    r_sub.italic = True
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Tabla de Datos Académicos y Equipo
    tbl_datos = doc.add_table(rows=7, cols=2)
    tbl_datos.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_datos.autofit = False

    datos_portada = [
        ("Asignatura:", "Estructura de Datos"),
        ("Cuatrimestre y Grupos:", "4° Cuatrimestre — Grupos 4A y 4D"),
        ("Profesora de la Asignatura:", "Ing. Mirian Magaly Canché Caamal, MGTI"),
        ("Integrantes del Equipo (3 personas):", 
         "1. __________________________________________________  Matrícula: ____________\n"
         "2. __________________________________________________  Matrícula: ____________\n"
         "3. __________________________________________________  Matrícula: ____________"),
        ("Fecha de Aplicación / Entrega:", "Septiembre – Octubre de 2026"),
        ("Puntaje / Calificación:", "100% (50% Saber + 50% Saber Hacer + Ser) — Meta: 10 de calificación"),
        ("Repositorio GitHub Oficial:", "https://github.com/Surer46/arraysexercises.git")
    ]

    for i, (label, val) in enumerate(datos_portada):
        cell_lbl = tbl_datos.cell(i, 0)
        cell_val = tbl_datos.cell(i, 1)
        
        cell_lbl.width = Inches(2.3)
        cell_val.width = Inches(4.2)
        
        set_cell_background(cell_lbl, "F8FAFC")
        set_cell_background(cell_val, "FFFFFF")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)
        
        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r_l = p0.add_run(label)
        r_l.bold = True
        r_l.font.name = "Segoe UI"
        r_l.font.size = Pt(9.5)
        r_l.font.color.rgb = RGBColor(15, 23, 42)
        
        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r_v = p1.add_run(val)
        r_v.font.name = "Segoe UI"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 2. INTRODUCCIÓN Y CONCEPTOS BÁSICOS EN LENGUAJE AMIGABLE
    # -------------------------------------------------------------
    add_heading_1(doc, "1. Introducción y Conceptos Básicos")
    add_paragraph(doc, 
        "En el mundo de la programación, una de las tareas más comunes es guardar y organizar muchos datos al mismo tiempo. "
        "Si quisiéramos guardar las calificaciones de 50 alumnos, sería muy pesado crear 50 variables diferentes (como alumno1, alumno2, etc.). "
        "Para solucionar este problema de forma elegante y rápida, existen los arreglos y las matrices.")

    add_heading_2(doc, "¿Qué es un Arreglo (Vector)?")
    add_paragraph(doc, 
        "Imagina un arreglo como una cajonera o casillero largo en la pared. Todos los cajones tienen el mismo nombre "
        "(por ejemplo, 'calificaciones'), pero cada cajón tiene un número para identificarlo. A ese número le llamamos índice. "
        "En la gran mayoría de los lenguajes modernos como C#, los cajones se empiezan a contar desde el número 0. "
        "Así, el primer dato está en el cajón [0], el segundo en el [1], y así sucesivamente.")

    add_heading_2(doc, "¿Qué es una Matriz (Arreglo de Dos Dimensiones)?")
    add_paragraph(doc, 
        "Una matriz es exactamente igual a una tabla de Excel o un tablero de ajedrez. "
        "En lugar de ser una sola fila de cajones, tiene filas (hacia abajo) y columnas (hacia la derecha). "
        "Para ubicar cualquier casilla dentro de la matriz, necesitamos dos números: el número de fila y el número de columna. "
        "Por ejemplo, matriz[0, 2] significa: 'mira la primera fila (0), en la tercera columna (2)'.")

    add_heading_2(doc, "¿Por qué utilizamos ciclos 'for'?")
    add_paragraph(doc, 
        "Para revisar todos los números de una matriz sin cansarnos, usamos dos ciclos 'for' (llamados ciclos anidados). "
        "El primer ciclo va pasando por cada fila de arriba hacia abajo, y el segundo ciclo va leyendo cada columna de izquierda a derecha. "
        "De esta manera, la computadora revisa cientos de datos en milisegundos sin equivocarse.")

    add_heading_2(doc, "Diseño Visual de las Aplicaciones (Estilo Cupertino y Blanco/Verde)")
    add_paragraph(doc, 
        "Para que las aplicaciones no parecieran ventanas aburridas o anticuadas, se aplicó un diseño moderno estilo Cupertino "
        "(inspirado en la limpieza y elegancia de Apple) con los siguientes principios visuales:")
    add_bullet(doc, "Tarjetas con bordes redondeados y sombras suaves para que los datos resalten a la vista.", "Estilo Cupertino:")
    add_bullet(doc, "Fondo blanco puro y gris perla, acompañado de un verde esmeralda fresco (#10B981) para botones y resultados.", "Paleta Blanco / Verde:")
    add_bullet(doc, "Iconos nítidos vectoriales de Material Design y tipografía clara Segoe UI con amplios espacios entre botones.", "Iconografía y Tipografía:")
    add_bullet(doc, "Si el usuario escribe letras en lugar de números o deja casillas vacías, el programa no se traba ni se cierra: muestra un aviso amable indicando qué corregir.", "Mensajes Claros de Validación:")

    # -------------------------------------------------------------
    # 3. LANZADOR GLOBAL: SUITE Y MENÚ PRINCIPAL
    # -------------------------------------------------------------
    add_heading_1(doc, "2. Menú Principal de la Suite (MenuPrincipal_Suite.exe)")
    add_paragraph(doc, 
        "El documento de evaluación solicita contar con un menú de opciones ordenado que permita abrir cada ejercicio "
        "sin necesidad de salir y volver a ejecutar todo el proyecto. "
        "Para cumplir esta meta, creamos el programa MenuPrincipal_Suite.exe. Al abrirlo, muestra una pantalla tipo tablero (dashboard) "
        "con una tarjeta elegante para cada uno de los 9 ejercicios.")

    add_image_with_caption(doc, 
        r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Evidencias_Capturas\captura_menu_principal.png", 
        "Menú Principal Cupertino con acceso a los 9 ejercicios y botón para abrir la carpeta de ejecutables.", width_inches=6.0)

    add_callout(doc, "💡 Características del Menú Suite:",
        "• Cada botón ejecuta directamente el archivo .exe correspondiente al ejercicio.\n"
        "• Puedes abrir varios ejercicios a la vez para comparar resultados sin que se cierre el menú principal.\n"
        "• Incluye un botón para abrir directamente la carpeta 'Ejecutables_EXE' en el Explorador de Archivos de Windows.")

    # -------------------------------------------------------------
    # 4. DESARROLLO DE LOS 9 EJERCICIOS (ANÁLISIS E-P-S + EVIDENCIAS)
    # -------------------------------------------------------------
    add_heading_1(doc, "3. Desarrollo Detallado de los Ejercicios (Valor: 80 Puntos)")
    add_paragraph(doc, 
        "A continuación se presenta el análisis completo de Entradas, Proceso y Salida (E-P-S), la explicación en palabras sencillas "
        "del método programado, la captura real del programa funcionando y las pruebas de validación de datos para cada ejercicio.")

    # --- EJERCICIO 1 ---
    add_heading_2(doc, "Ejercicio 1: Contador de Ceros por Renglón")
    add_paragraph(doc, "Elabore un programa en C# que calcule cuántos ceros aparecen en cada renglón de un arreglo de 5x5 números.", "Descripción del Problema: ")

    add_heading_3(doc, "Análisis de Entradas, Proceso y Salida (E-P-S):")
    add_bullet(doc, "Una matriz cuadrada de 5 filas y 5 columnas de números enteros (la del PDF o valores que escriba el usuario).", "Entradas:")
    add_bullet(doc, "Se lee renglón por renglón. Para cada renglón, se cuenta cuántas veces aparece el número 0. Este conteo se guarda en una lista o arreglo.", "Proceso:")
    add_bullet(doc, "La cantidad de ceros de cada renglón (ej. Renglón 1: 1 cero, Renglón 2: 3 ceros, etc.) y la suma total de ceros en toda la matriz.", "Salida:")

    add_heading_3(doc, "Explicación del Método Separado (MatrizContadorCerosService):")
    add_paragraph(doc, 
        "Creamos una función especial llamada ContarCerosPorRenglon(matriz). "
        "Esta función crea un arreglo de 5 números para guardar los resultados. "
        "Luego, con un ciclo 'for' viaja fila por fila. En cada fila pone un contador en 0 y revisa las 5 columnas. "
        "Si encuentra una celda con valor 0, le suma 1 al contador. Al terminar la fila, guarda ese total y pasa a la siguiente fila. "
        "En la pantalla, las casillas que tienen un cero se pintan automáticamente de color verde menta para que se identifiquen al instante.")

    add_image_with_caption(doc, 
        r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Evidencias_Capturas\captura_ejercicio1.png", 
        "Ejercicio 1 funcionando con la matriz del PDF y ceros resaltados.", width_inches=5.8)

    add_callout(doc, "🛡️ Validación de Errores:",
        "Si el usuario borra un número y deja la casilla vacía, el programa avisa: 'Hay datos faltantes en la fila X, columna Y'. "
        "Si escribe letras en vez de números, el sistema responde: 'Introducir sólo números enteros'.")

    # --- EJERCICIO 2 ---
    add_heading_2(doc, "Ejercicio 2: Cuadrado Mágico y Constante Mágica")
    add_paragraph(doc, 
        "Un cuadrado mágico es una cuadrícula de números enteros donde la suma de cada fila, de cada columna y de las dos diagonales principales "
        "da exactamente el mismo resultado, llamado constante mágica. El programa debe leer la matriz y decir si es o no un cuadrado mágico.", 
        "Descripción del Problema: ")

    add_heading_3(doc, "Análisis de Entradas, Proceso y Salida (E-P-S):")
    add_bullet(doc, "El tamaño del cuadrado (por ejemplo 3x3, 4x4 o 5x5) y los números enteros colocados en cada casilla.", "Entradas:")
    add_bullet(doc, "1) Sumar los elementos de cada fila. 2) Sumar los elementos de cada columna. 3) Sumar la diagonal principal y la secundaria. 4) Comparar si todas las sumas son iguales.", "Proceso:")
    add_bullet(doc, "Un mensaje que dice claramente si es o no Cuadrado Mágico, mostrando el valor de la Constante Mágica y las sumas en los bordes.", "Salida:")

    add_heading_3(doc, "Explicación del Método Separado (CuadroMagicoService):")
    add_paragraph(doc, 
        "La función AnalizarCuadroMagico(matriz) primero toma la suma de la primera fila como 'meta'. "
        "Luego recorre las demás filas y columnas sumando sus valores. Para la diagonal principal suma las casillas donde la fila y la columna son iguales [i, i]. "
        "Para la diagonal secundaria suma las casillas de la otra esquina [i, n - 1 - i]. "
        "Si todas las sumas dan el mismo número, devuelve '¡Es Cuadrado Mágico!' junto con la constante (en el ejemplo del PDF la constante es 15). "
        "Además, incluye un botón para generar automáticamente cuadrados mágicos perfectos usando el método siamés.")

    add_image_with_caption(doc, 
        r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Evidencias_Capturas\captura_ejercicio2.png", 
        "Ejercicio 2 verificando el cuadrado mágico 3x3 con constante mágica 15.", width_inches=5.8)

    # --- EJERCICIO 3 ---
    add_heading_2(doc, "Ejercicio 3: Operaciones con Matrices 2x2")
    add_paragraph(doc, 
        "Crear una aplicación que lea dos matrices de tamaño 2x2 y calcule: a) La suma, b) La resta de la primera menos la segunda, "
        "c) El producto simple elemento por elemento, d) La división simple elemento por elemento.", 
        "Descripción del Problema: ")

    add_heading_3(doc, "Análisis de Entradas, Proceso y Salida (E-P-S):")
    add_bullet(doc, "Dos matrices de 2 filas y 2 columnas con números (Matriz 1 y Matriz 2).", "Entradas:")
    add_bullet(doc, "Para cada casilla [i, j], se realiza la operación matemática directa entre el número de la Matriz 1 y el de la Matriz 2.", "Proceso:")
    add_bullet(doc, "Cuatro cuadrículas con los resultados: Suma, Resta, Multiplicación y División (redondeada a 2 decimales).", "Salida:")

    add_heading_3(doc, "Explicación de los Métodos Separados (Matrices2x2Service):")
    add_paragraph(doc, 
        "Separamos cada operación en su propio método: Sumar, Restar, ProductoSimple y DivisionSimple. "
        "Cada método usa dos ciclos for sencillos que van del 0 al 1. Por ejemplo, en la suma hace: Resultado[i, j] = MatrizA[i, j] + MatrizB[i, j]. "
        "En la división, antes de hacer el cálculo se revisa con un 'if' que el número de la Matriz 2 no sea cero. "
        "Si alguien pone un 0 como divisor, el programa lo detecta amablemente y avisa que no se puede dividir entre cero.")

    add_image_with_caption(doc, 
        r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Evidencias_Capturas\captura_ejercicio3.png", 
        "Ejercicio 3 mostrando los resultados de las 4 operaciones simultáneas con los datos del PDF.", width_inches=5.8)

    # --- EJERCICIO 4 ---
    add_heading_2(doc, "Ejercicio 4: Generador de Matriz Identidad")
    add_paragraph(doc, 
        "Desarrollar una aplicación para llenar una matriz cuadrada de números enteros de manera que: "
        "a) La diagonal principal contenga 1s (unos) y b) Las demás posiciones contengan 0s (ceros).", 
        "Descripción del Problema: ")

    add_heading_3(doc, "Análisis de Entradas, Proceso y Salida (E-P-S):")
    add_bullet(doc, "El tamaño deseado N para la matriz (por ejemplo 3, 5, 8 o 10).", "Entradas:")
    add_bullet(doc, "Un ciclo doble donde si el número de fila es igual al número de columna (i == j), se pone un 1. En caso contrario, se pone un 0.", "Proceso:")
    add_bullet(doc, "Una matriz visualmente hermosa donde los 1s brillan en verde esmeralda en diagonal y los 0s quedan en tono suave.", "Salida:")

    add_heading_3(doc, "Explicación del Método Separado (MatrizIdentidadService):")
    add_paragraph(doc, 
        "El método GenerarMatrizIdentidad(n) crea una matriz de tamaño n x n en la memoria. "
        "Recorre todas las casillas con dos ciclos for: 'si i == j entonces matriz[i,j] = 1, si no matriz[i,j] = 0'. "
        "En matemáticas y en programación gráfica, la matriz identidad es súper importante porque funciona como el 'número 1' de las matrices: "
        "cualquier matriz multiplicada por la identidad se queda exactamente igual.")

    add_image_with_caption(doc, 
        r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Evidencias_Capturas\captura_ejercicio4.png", 
        "Ejercicio 4 generando una matriz identidad dinámica con diagonal resaltada.", width_inches=5.8)

    # --- EJERCICIO 5 ---
    add_heading_2(doc, "Ejercicio 5: Matriz 5x10 con Sumas y Promedios (Arreglos A, B, C, D)")
    add_paragraph(doc, 
        "Llenar una matriz de 5 filas y 10 columnas con números y calcular la suma y promedio de cada fila y de cada columna, "
        "guardándolos en 4 arreglos independientes: A (suma fila), B (promedio fila), C (suma columna) y D (promedio columna).", 
        "Descripción del Problema: ")

    add_heading_3(doc, "Análisis de Entradas, Proceso y Salida (E-P-S):")
    add_bullet(doc, "Una matriz de 50 casillas (5 renglones por 10 columnas) con números enteros.", "Entradas:")
    add_bullet(doc, "1) Sumar horizontalmente cada una de las 5 filas y dividir entre 10 para su promedio. "
                    "2) Sumar verticalmente cada una de las 10 columnas y dividir entre 5 para su promedio.", "Proceso:")
    add_bullet(doc, "La tabla central 5x10 con los arreglos A y B a la derecha, y los arreglos C y D en la parte inferior, igual a la imagen del PDF.", "Salida:")

    add_heading_3(doc, "Explicación del Método Separado (EstadisticasMatrizService):")
    add_paragraph(doc, 
        "La función CalcularEstadisticas(matriz) crea 4 arreglos independientes. "
        "Primero hace un recorrido por filas para llenar el Arreglo A (suma de los 10 valores de la fila) y el Arreglo B (suma dividida entre 10). "
        "Luego hace un recorrido por columnas para llenar el Arreglo C (suma de los 5 valores de la columna) y el Arreglo D (suma dividida entre 5). "
        "El diseño respeta fielmente el formato solicitado en la rúbrica del instrumento.")

    add_image_with_caption(doc, 
        r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Evidencias_Capturas\captura_ejercicio5.png", 
        "Ejercicio 5 con el formato oficial de la matriz 5x10 y vectores laterales e inferiores.", width_inches=5.8)

    # --- EJERCICIO 6 ---
    add_heading_2(doc, "Ejercicio 6: Resumen Anual de Ventas (12 Meses x 5 Días)")
    add_paragraph(doc, 
        "Dada una tabla de ventas donde las 12 filas son los meses del año (Enero a Diciembre) y las 5 columnas son días de la semana "
        "(Lunes a Viernes), calcular: menor venta con su mes y día, mayor venta con su mes y día, venta total y venta acumulada por día.", 
        "Descripción del Problema: ")

    add_heading_3(doc, "Análisis de Entradas, Proceso y Salida (E-P-S):")
    add_bullet(doc, "La matriz de ventas de 12 renglones por 5 columnas con los datos de la práctica.", "Entradas:")
    add_bullet(doc, "Se recorre toda la tabla comparando cada número para encontrar el más pequeño y el más grande guardando su mes y día. "
                    "Se acumula la suma general y se suma cada columna para saber cuánto se vendió cada día de la semana.", "Proceso:")
    add_bullet(doc, "Menor venta ($3.00 en Noviembre, Martes), Mayor venta ($82.00 en Noviembre, Viernes), Venta total ($2,056.00) y venta por día.", "Salida:")

    add_heading_3(doc, "Explicación del Método Separado (ResumenVentasService):")
    add_paragraph(doc, 
        "El método AnalizarVentas(matriz) empieza asumiendo que la venta de Enero-Lunes es la menor y la mayor. "
        "Luego va revisando celda por celda: si encuentra un valor más chico, actualiza el valor mínimo y anota el nombre del mes y del día. "
        "Si encuentra uno más grande, hace lo mismo para el máximo. "
        "Al mismo tiempo va sumando el total general y sumando en una bolsa para cada día de la semana. "
        "En la tabla visual, la venta más baja se pinta de amarillo suave y la más alta de verde esmeralda brillante.")

    add_image_with_caption(doc, 
        r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Evidencias_Capturas\captura_ejercicio6.png", 
        "Ejercicio 6 con el reporte comercial completo y celdas récord resaltadas.", width_inches=5.8)

    # --- EJERCICIO 7 ---
    add_heading_2(doc, "Ejercicio 7: Resumen y Análisis de Calificaciones Escolares")
    add_paragraph(doc, 
        "Dada una tabla de calificaciones de 12 alumnos (filas) y 3 parciales (columnas), calcular: "
        "el promedio de cada alumno, el promedio más alto, el más bajo, cuántos parciales fueron reprobados (< 7.0) "
        "y la tabla de distribución de calificaciones finales en 6 rangos de notas.", 
        "Descripción del Problema: ")

    add_heading_3(doc, "Análisis de Entradas, Proceso y Salida (E-P-S):")
    add_bullet(doc, "Arreglo de 12 filas y 3 columnas con notas decimales de 0 a 10.", "Entradas:")
    add_bullet(doc, "Promediar los 3 parciales de cada estudiante, encontrar el mayor y menor promedio, contar cuántas notas individuales son menores a 7.0, "
                    "y clasificar los promedios en 6 intervalos de notas.", "Proceso:")
    add_bullet(doc, "Promedios individuales con distintivo (Aprobado/Reprobado), promedio más alto (9.40), más bajo (5.00), total reprobados (13) y la tabla de frecuencias.", "Salida:")

    add_heading_3(doc, "Explicación del Método Separado (ResumenCalificacionesService):")
    add_paragraph(doc, 
        "El método AnalizarCalificaciones(matriz) calcula primero el promedio sumando los 3 parciales y dividiendo entre 3. "
        "Durante el recorrido, si una nota de parcial es menor a 7.0, incrementa un contador de reprobados. "
        "Después clasifica cada promedio final en 6 grupos: [0 a 4.9], [5.0 a 5.9], [6.0 a 6.9], [7.0 a 7.9], [8.0 a 8.9] y [9.0 a 10.0]. "
        "Esto genera una gráfica y lista idéntica a la solicitada en el documento oficial.")

    add_image_with_caption(doc, 
        r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Evidencias_Capturas\captura_ejercicio7.png", 
        "Ejercicio 7 con análisis académico, estado de aprobación y distribución de frecuencias.", width_inches=5.8)

    # --- EJERCICIO 8 (NIVEL ALTO 10/10) ---
    add_heading_2(doc, "Ejercicio 8 (Nivel de Dificultad Alto): Búsqueda de Camino Óptimo en Matriz 2D")
    add_paragraph(doc, 
        "La rúbrica indica que para alcanzar el 10 de calificación se deben agregar 2 ejercicios de dificultad alta con arreglos. "
        "En este ejercicio se implementa la búsqueda del camino más corto en una cuadrícula (laberinto con obstáculos) desde un punto de inicio [A] "
        "hasta un punto de destino [B] utilizando arreglos bidimensionales y el algoritmo BFS (búsqueda en anchura).", 
        "Descripción del Problema y Justificación: ")

    add_heading_3(doc, "Análisis de Entradas, Proceso y Salida (E-P-S):")
    add_bullet(doc, "Una matriz interactiva de 8x8 donde el usuario puede dar clic en cualquier casilla para colocar o quitar muros/obstáculos.", "Entradas:")
    add_bullet(doc, "Se usa una matriz de visitados (bool[,]) y una matriz de padres ((int,int)[,]) para explorar casilla por casilla en las 4 direcciones "
                    "(arriba, abajo, izquierda, derecha) sin pasar por muros hasta llegar a la meta.", "Proceso:")
    add_bullet(doc, "El camino óptimo trazado con celdas verde esmeralda numeradas paso a paso, la cantidad de nodos explorados y la lista de coordenadas.", "Salida:")

    add_heading_3(doc, "Explicación del Método Separado (LaberintoService):")
    add_paragraph(doc, 
        "El método EncontrarRutaOptima(mapa, inicio, destino) funciona como un GPS en miniatura. "
        "Pone la casilla de inicio en una fila de espera (cola) y la marca en una matriz booleana para no visitarla dos veces. "
        "Luego saca la casilla y mira sus 4 vecinas. Si una vecina está libre y dentro de la matriz, anota de dónde vino y la mete a la fila. "
        "Cuando llega a la meta, viaja hacia atrás siguiendo las anotaciones de padres y reconstruye el camino más corto posible.")

    add_image_with_caption(doc, 
        r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Evidencias_Capturas\captura_ejercicio8.png", 
        "Ejercicio 8 encontrando la ruta más corta paso a paso sorteando muros interactivos.", width_inches=5.8)

    # --- EJERCICIO 9 (NIVEL ALTO 10/10) ---
    add_heading_2(doc, "Ejercicio 9 (Nivel de Dificultad Alto): Multiplicación Matricial Avanzada")
    add_paragraph(doc, 
        "Segundo ejercicio de dificultad alta. Implementa la multiplicación completa de dos matrices de dimensiones arbitrarias "
        "(Matriz A de NxM por Matriz B de MxP), el cálculo de la matriz transpuesta y el cálculo del determinante por cofactores recursivos.", 
        "Descripción del Problema y Justificación: ")

    add_heading_3(doc, "Análisis de Entradas, Proceso y Salida (E-P-S):")
    add_bullet(doc, "Dimensiones N, M y P, y los valores de las matrices A y B (con botones de dimensiones preestablecidas).", "Entradas:")
    add_bullet(doc, "1) Se valida que las columnas de A coincidan con las filas de B. 2) Se ejecuta el triple ciclo for para calcular la sumatoria de productos. "
                    "3) Se invierten filas por columnas para la transpuesta. 4) Se calcula el determinante mediante submatrices.", "Proceso:")
    add_bullet(doc, "La matriz producto C resultante, la matriz transpuesta, el determinante y el desglose de fórmulas paso a paso celda por celda.", "Salida:")

    add_heading_3(doc, "Explicación del Método Separado (AlgebraMatricialService):")
    add_paragraph(doc, 
        "El método MultiplicarMatrices utiliza tres ciclos for: el primero recorre las filas de A, el segundo las columnas de B, "
        "y el tercer ciclo va multiplicando parejas de números y sumándolos: C[i, j] += A[i, k] * B[k, j]. "
        "En pantalla se genera una explicación paso a paso que le muestra al alumno exactamente de dónde salió cada número. "
        "Además, la función Transponer intercambia las coordenadas para convertir filas en columnas de manera instantánea.")

    add_image_with_caption(doc, 
        r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Evidencias_Capturas\captura_ejercicio9.png", 
        "Ejercicio 9 realizando la multiplicación matricial completa con fórmulas detalladas y determinante.", width_inches=5.8)

    # -------------------------------------------------------------
    # 5. RESUMEN DE PRUEBAS Y VALIDACIÓN DE EXCEPCIONES
    # -------------------------------------------------------------
    add_heading_1(doc, "4. Pruebas de Validación y Manejo de Excepciones")
    add_paragraph(doc, 
        "Para garantizar que los programas sean estables y confiables (criterio de 30 puntos en la rúbrica), "
        "se implementaron validaciones que protegen al usuario ante cualquier error de captura:")

    tbl_pruebas = doc.add_table(rows=6, cols=3)
    tbl_pruebas.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pruebas.autofit = False

    headers_pruebas = ["Situación Probada", "Entrada del Usuario", "Comportamiento del Sistema y Mensaje"]
    for col_idx, h_text in enumerate(headers_pruebas):
        cell = tbl_pruebas.cell(0, col_idx)
        set_cell_background(cell, "10B981")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    casos_pruebas = [
        ("Casilla Vacía", "El usuario borra un número y deja el campo en blanco.", "Muestra advertencia: 'Hay datos faltantes en la fila X, columna Y. Por favor completa la celda.'"),
        ("Texto no Numérico", "El usuario escribe 'abc' o caracteres especiales.", "Muestra advertencia: 'El valor no es válido. Introducir sólo números.'"),
        ("División por Cero", "En la Matriz 2 del Ejercicio 3 se coloca un valor 0.", "Evita que el programa se cierre y avisa: 'División por cero en posición [i,j]: B es 0.'"),
        ("Nota fuera de Rango", "En el Ejercicio 7 se introduce una calificación de 15 o -2.", "Valida que las notas estén entre 0 y 10: 'La calificación debe estar en el rango de 0 a 10.'"),
        ("Matriz no Compatible", "En multiplicación matricial si las dimensiones no coinciden.", "Avisa que las columnas de A deben ser iguales a las filas de B para poder multiplicarse.")
    ]

    for row_idx, (sit, ent, comp) in enumerate(casos_pruebas, start=1):
        for col_idx, val in enumerate([sit, ent, comp]):
            cell = tbl_pruebas.cell(row_idx, col_idx)
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Segoe UI"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)

    add_paragraph(doc, "", space_after=8)

    # -------------------------------------------------------------
    # 6. CONCLUSIONES Y RECOMENDACIONES
    # -------------------------------------------------------------
    add_heading_1(doc, "5. Conclusiones y Aprendizajes")
    add_paragraph(doc, 
        "El desarrollo de esta práctica permitió consolidar los conceptos fundamentales de estructuras de datos en C#:")
    add_bullet(doc, "Comprendimos que los arreglos y matrices son la base sobre la cual se construyen estructuras más avanzadas como listas, pilas, colas y grafos.", "Importancia de los Arreglos:")
    add_bullet(doc, "Separar la lógica en clases dentro de 'Servicios/' permite que el código sea limpio, fácil de leer y muy sencillo de probar y mantener.", "Buenas Prácticas de Programación (POO):")
    add_bullet(doc, "Un buen programa no solo debe calcular números correctamente, sino también tener una interfaz amigable, intuitiva y visualmente atractiva para el usuario final.", "Experiencia de Usuario (UI/UX):")
    add_bullet(doc, "Todos los archivos del proyecto, incluyendo el código fuente, la solución y los ejecutables independientes, están disponibles de forma pública y respaldados en GitHub.", "Entorno de Trabajo:")

    # Guardar documento
    output_docx = r"C:\Universidad\Cuatrimestre_4\Estructura de datos\EjerciciosArrays\Practica_Arreglos_Estructura_de_Datos_UTM.docx"
    doc.save(output_docx)
    print(f"Documento generado exitosamente en: {output_docx}")

if __name__ == "__main__":
    create_document()
